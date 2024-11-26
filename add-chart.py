from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

# Load the workbook and select the worksheet
wb = load_workbook('requirement/pivot_table.xlsx')
sheet = wb['report']

# Create a BarChart object
barchart = BarChart()

# Define the data range for the bar chart
data = Reference(sheet, min_col=2, min_row=2, max_col=13, max_row=15)

# Define the categories (x-axis labels) for the bar chart
categories = Reference(sheet, min_col=1, min_row=1, max_col=1, max_row=15)

# Add data to the bar chart, specifying that the titles are from the first row of data
barchart.add_data(data, titles_from_data=True)

# Set the categories for the bar chart
barchart.set_categories(categories)

# Add the bar chart to the worksheet at the specified position
sheet.add_chart(barchart, 'O2')

# Optionally set the title and style for the bar chart
# barchart.title = "salesDetails"
# barchart.style = 5

# Save the workbook with the new bar chart
wb.save('requirement/barchart.xlsx')
